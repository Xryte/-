from datetime import timedelta

from django.contrib.admin.views.decorators import staff_member_required
from django.db.models import Count
from django.http import HttpResponse
from django.shortcuts import render
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

from ads.models import Ad
from .models import PageView


@staff_member_required
def dashboard(request):
    now = timezone.now()
    period = request.GET.get('period', 'week')

    period_map = {
        'day': timedelta(days=1),
        'week': timedelta(weeks=1),
        'month': timedelta(days=30),
        'year': timedelta(days=365),
        'all': None,
    }
    delta = period_map.get(period)

    visits_qs = PageView.objects.all()
    if delta:
        visits_qs = visits_qs.filter(visited_at__gte=now - delta)

    total_visits = visits_qs.count()
    unique_ips = visits_qs.values('ip_address').distinct().count()

    top_pages = (
        visits_qs
        .values('path')
        .annotate(count=Count('id'))
        .order_by('-count')[:10]
    )

    top_ads = Ad.objects.order_by('-views_count')[:10]

    if delta:
        visits_by_day = (
            visits_qs
            .extra(select={'day': "DATE(visited_at)"})
            .values('day')
            .annotate(count=Count('id'))
            .order_by('day')
        )
    else:
        visits_by_day = []

    context = {
        'total_visits': total_visits,
        'unique_ips': unique_ips,
        'top_pages': top_pages,
        'top_ads': top_ads,
        'visits_by_day': list(visits_by_day),
        'current_period': period,
        'periods': period_map.keys(),
    }
    return render(request, 'analytics/dashboard.html', context)


@staff_member_required
def export_excel(request):
    period = request.GET.get('period', 'month')
    period_map = {
        'day': timedelta(days=1),
        'week': timedelta(weeks=1),
        'month': timedelta(days=30),
        'year': timedelta(days=365),
        'all': None,
    }
    delta = period_map.get(period)

    visits_qs = PageView.objects.select_related('user').all()
    if delta:
        visits_qs = visits_qs.filter(visited_at__gte=timezone.now() - delta)

    wb = Workbook()
    ws = wb.active
    ws.title = 'Посещения'

    headers = ['ID', 'Путь', 'Метод', 'IP-адрес', 'User-Agent', 'Пользователь', 'Время посещения']
    header_font = Font(bold=True, color='FFFFFF', size=11)
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')

    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')

    for row_num, visit in enumerate(visits_qs.iterator(), start=2):
        ws.cell(row=row_num, column=1, value=visit.id)
        ws.cell(row=row_num, column=2, value=visit.path)
        ws.cell(row=row_num, column=3, value=visit.method)
        ws.cell(row=row_num, column=4, value=visit.ip_address or '')
        ws.cell(row=row_num, column=5, value=visit.user_agent)
        ws.cell(row=row_num, column=6, value=visit.user.username if visit.user else 'Аноним')
        ws.cell(row=row_num, column=7, value=visit.visited_at.strftime('%Y-%m-%d %H:%M:%S'))

    widths = [8, 40, 8, 18, 50, 20, 22]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = width

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="visits_{period}.xlsx"'
    wb.save(response)
    return response